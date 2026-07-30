"""Extract plain text from uploaded CV files (PDF / DOCX / MD / TXT)."""

from __future__ import annotations

from pathlib import PurePosixPath

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt"}
REJECTED_EXTENSIONS = {".doc"}


class CvIngestError(ValueError):
    """Unsupported or unreadable CV upload."""


def normalize_filename(filename: str | None) -> str:
    name = (filename or "upload.bin").strip() or "upload.bin"
    return PurePosixPath(name).name


def extension_of(filename: str | None) -> str:
    return PurePosixPath(normalize_filename(filename)).suffix.lower()


def extract_text(filename: str | None, data: bytes) -> str:
    """Return UTF-8 text extracted from upload bytes."""
    if not data:
        raise CvIngestError("Leere Datei")
    ext = extension_of(filename)
    if ext in REJECTED_EXTENSIONS:
        raise CvIngestError(
            "Legacy .doc wird nicht unterstützt. Bitte als PDF oder DOCX speichern."
        )
    if ext not in SUPPORTED_EXTENSIONS:
        raise CvIngestError(
            f"Format {ext or '(ohne Endung)'} nicht unterstützt. "
            "Erlaubt: PDF, DOCX, MD, TXT."
        )
    if ext == ".pdf":
        return _extract_pdf(data)
    if ext == ".docx":
        return _extract_docx(data)
    # md / markdown / txt
    return _decode_text(data)


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise CvIngestError("Textdatei nicht dekodierbar")
    text = text.strip()
    if not text:
        raise CvIngestError("Keine Textinhalte erkannt")
    return text


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise CvIngestError("pypdf ist nicht installiert") from exc
    from io import BytesIO

    reader = PdfReader(BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        chunk = page.extract_text() or ""
        if chunk.strip():
            parts.append(chunk.strip())
    text = "\n\n".join(parts).strip()
    if not text:
        raise CvIngestError("PDF enthält keinen extrahierbaren Text")
    return text


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise CvIngestError("python-docx ist nicht installiert") from exc
    from io import BytesIO

    doc = Document(BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    # tables
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise CvIngestError("DOCX enthält keinen extrahierbaren Text")
    return text
